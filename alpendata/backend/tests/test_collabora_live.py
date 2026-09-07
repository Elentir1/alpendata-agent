"""Real CODE + Chromium + AlpenData UI + TLS + PostgreSQL; optional infrastructure recipe."""

import base64
import io
import socket
import subprocess
import sys
import threading
import time
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
import requests
import uvicorn
from fastapi.staticfiles import StaticFiles
from sqlalchemy import select
from starlette.middleware.base import BaseHTTPMiddleware
from test_ingress import certificate, wait_http

from alpendata_api.auth import SESSION_COOKIE
from alpendata_api.file_store import FileStoreSettings
from alpendata_api.ingress import content_security_policy
from alpendata_api.model_gateway import ModelSettings
from alpendata_api.models import EditorSession
from alpendata_api.runtime import RuntimeSettings
from alpendata_api.settings import Settings

pytestmark = pytest.mark.linux_only


def port():
    with socket.socket() as reserve:
        reserve.bind(("127.0.0.1", 0))
        return reserve.getsockname()[1]


@pytest.fixture
def service(request, database_url, service_factory, tmp_path, monkeypatch):
    image, browser, dist = (
        request.config.getoption(name) for name in ("--collabora-image", "--chromium-bin", "--frontend-dist")
    )
    if not image or not browser or not dist:
        pytest.skip("Supply --collabora-image, --chromium-bin and --frontend-dist for real CODE recipe")
    assert image.startswith("sha256:") and len(image) == 71
    output_root = request.config.getoption("--document-qa-output")
    output = Path(output_root) / request.node.callspec.id if output_root else tmp_path
    output.mkdir(mode=0o700, parents=True, exist_ok=True)
    api_port, office_port = port(), port()
    origin, office = f"https://127.0.0.1:{api_port}", f"https://127.0.0.1:{office_port}"
    tls = tmp_path / "tls"
    tls.mkdir(mode=0o755)
    cert, key = certificate(tls)
    key.chmod(0o644)  # Disposable test key, readable by container UID; no production credentials.

    class TrustedSession(requests.Session):
        def __init__(self):
            super().__init__()
            self.verify = str(cert)

    monkeypatch.setattr("alpendata_api.office_discovery.requests.Session", TrustedSession)
    name = "alpendata-code-qa-" + uuid4().hex[:10]
    args = [
        "podman",
        "run",
        "--detach",
        "--name",
        name,
        "--network=host",
        "--cap-drop=ALL",
        "--security-opt=no-new-privileges",
        "--memory=2g",
        "--cpus=2",
        "--pids-limit=256",
        "--volume",
        str(tls) + ":/qa:ro",
        image,
        "--port=" + str(office_port),
        "--o:net.listen=loopback",
        "--o:net.proto=IPv4",
        "--o:security.capabilities=false",
        "--o:ssl.enable=true",
        "--o:ssl.cert_file_path=/qa/test.pem",
        "--o:ssl.key_file_path=/qa/test.key",
        "--o:ssl.ca_file_path=/qa/test.pem",
        "--o:storage.ssl.ca_file_path=/qa/test.pem",
        "--o:ssl.ssl_verification=true",
        "--o:server_name=127.0.0.1:" + str(office_port),
        "--o:storage.wopi.alias_groups[@mode]=groups",
        "--o:storage.wopi.alias_groups.group[0].host=" + origin,
        "--o:storage.wopi.alias_groups.group[0].host[@allow]=true",
        "--o:admin_console.enable=false",
        "--o:logging.level=warning",
        "--o:ai.enabled=false",
        "--o:ai.allow_user_settings=false",
        "--o:zotero.enable=false",
        "--o:security.enable_macros_execution=false",
        "--o:num_prespawn_children=1",
    ]
    subprocess.run(args, check=True, capture_output=True, timeout=30)
    try:
        with httpx.Client(verify=str(cert), trust_env=False, base_url=office) as check:
            wait_http(check, "/hosting/discovery", lambda: True)
        settings = Settings(
            database_url=database_url,
            public_origin=origin,
            office_origin=office,
            office_secret="synthetic-collabora-test-signing-secret",
            files=FileStoreSettings(root=tmp_path / "objects"),
            workspace_organizations=("*",),
            model=ModelSettings("mistral", "synthetic-model", "synthetic-key"),
            runtime=RuntimeSettings(tmp_path / "states", "sha256:" + "0" * 64, executable=sys.executable),
            smtp_host="smtp.example.test",
            smtp_sender="noreply@example.com",
            smtp_username="synthetic",
            smtp_password="synthetic",
        )
        with service_factory(settings) as pair:
            app, _ = pair
            app.mount("/", StaticFiles(directory=dist, html=True))

            async def policy(request, call_next):
                response = await call_next(request)
                response.headers["Content-Security-Policy"] = content_security_policy(office)
                return response

            server = uvicorn.Server(
                uvicorn.Config(
                    BaseHTTPMiddleware(app, dispatch=policy),
                    host="127.0.0.1",
                    port=api_port,
                    ssl_certfile=str(cert),
                    ssl_keyfile=str(key),
                    access_log=False,
                    log_level="critical",
                )
            )
            thread = threading.Thread(target=server.run, daemon=True)
            thread.start()
            try:
                with httpx.Client(verify=str(cert), trust_env=False, base_url=origin) as check:
                    wait_http(check, "/health/ready", thread.is_alive)
                app.state.office_qa = {
                    "browser": browser,
                    "origin": origin,
                    "office": office,
                    "output": output,
                }
                yield pair
            finally:
                server.should_exit = True
                thread.join(timeout=15)
    finally:
        result = subprocess.run(["podman", "logs", "--tail", "100", name], capture_output=True, timeout=15)
        # Tokens may occur in provider logs; retain privately on QA, never print them.
        (output / "collabora.log").write_bytes(result.stdout + result.stderr)
        subprocess.run(["podman", "rm", "--force", name], check=True, capture_output=True, timeout=30)


def office_bytes(extension, text):
    buffer = io.BytesIO()
    if extension == "docx":
        from docx import Document

        document = Document()
        document.add_paragraph(text)
    elif extension == "xlsx":
        from openpyxl import Workbook

        document = Workbook()
        document.active["A1"] = text
    else:
        from pptx import Presentation

        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text = text
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "language,viewport",
    [("en", {"width": 1600, "height": 1050}), ("fr", {"width": 390, "height": 844})],
    ids=["desktop-en", "mobile-fr"],
)
def test_real_editor_opens_office_formats_saves_selection_and_preserves_conflict(
    service, account, language, viewport
):
    from playwright.sync_api import expect, sync_playwright

    app, client = service
    _, owner = account("office-qa@example.com")
    org = client.post("/api/organizations", headers=owner, json={"name": "Office QA"}).json()["id"]
    root = f"/api/organizations/{org}"
    client.put(
        root + "/onboarding",
        headers=owner,
        json={"language": language, "role": "Coach", "needs": "Documents"},
    )
    chat = client.post(root + "/chat/conversations", headers=owner, json={}).json()["id"]
    files = {}
    for extension in ("docx", "xlsx", "pptx"):
        result = client.post(
            root + "/chat/conversations/" + chat + "/files",
            headers=owner,
            json={
                "request_id": str(uuid4()),
                "filename": "Pilot." + extension,
                "content_base64": base64.b64encode(
                    office_bytes(extension, "AlpenData synthetic pilot")
                ).decode(),
            },
        )
        assert result.status_code == 201, result.text
        files[extension] = result.json()["id"]
    qa = app.state.office_qa
    with sync_playwright() as engine:
        browser = engine.chromium.launch(executable_path=qa["browser"], headless=True)
        context = browser.new_context(
            ignore_https_errors=True,
            viewport=viewport,
            is_mobile=language == "fr",
            has_touch=language == "fr",
            **(
                {
                    "user_agent": (
                        "Mozilla/5.0 (Linux; Android 14) AppleWebKit/537.36 "
                        "(KHTML, like Gecko) Chrome/152.0.0.0 Mobile Safari/537.36"
                    )
                }
                if language == "fr"
                else {}
            ),
        )
        context.add_cookies(
            [
                {
                    "name": SESSION_COOKIE,
                    "value": owner["Authorization"].removeprefix("Bearer "),
                    "url": qa["origin"],
                    "secure": True,
                    "httpOnly": True,
                    "sameSite": "Lax",
                }
            ]
        )
        context.add_init_script(f"localStorage.setItem('alpendata.language', '{language}')")
        page = context.new_page()
        welcome = page.frame_locator("iframe").locator(".iframe-welcome-wrap")
        page.add_locator_handler(
            welcome, lambda: page.frame_locator("iframe").locator("body").press("Escape")
        )
        try:
            for extension in ("docx", "xlsx", "pptx"):
                page.goto(qa["origin"] + "/?organization=" + org + "&conversation=" + chat)
                page.get_by_role(
                    "button", name="Fichiers" if language == "fr" else "Files", exact=True
                ).click()
                page.locator(".workspace-file-list button").filter(has_text="Pilot." + extension).click()
                page.get_by_role(
                    "button", name="Ouvrir l’éditeur" if language == "fr" else "Open editor", exact=True
                ).click()
                save = page.get_by_role(
                    "button",
                    name="Enregistrer le document" if language == "fr" else "Save document",
                    exact=True,
                )
                expect(save).to_be_enabled(timeout=60000)
                if welcome.is_visible():
                    page.frame_locator("iframe").locator("body").press("Escape")
                    expect(welcome).to_be_hidden()
                if language == "fr":
                    edit_button = page.frame_locator("iframe").locator("#mobile-edit-button")
                    edit_button.tap()
                    expect(edit_button).to_be_hidden()
                assert page.locator("iframe").bounding_box()["y"] < viewport["height"] / 2
                page.screenshot(path=str(qa["output"] / (extension + "-opened.png")))
                (qa["output"] / (extension + "-editor-ui.txt")).write_text(
                    page.frame_locator("iframe").locator("body").inner_text()
                )
                if extension != "docx":
                    # Force a real round-trip save from Calc and Impress as well.
                    save.click()
                    file = root + "/files/" + files[extension]
                    deadline = time.monotonic() + 30
                    while time.monotonic() < deadline:
                        versions = client.get(file + "/versions", headers=owner).json()
                        if versions["file"]["version"] >= 2:
                            break
                        time.sleep(0.2)
                    assert versions["file"]["version"] == 2
                    saved = client.get(file + "/versions/2/content", headers=owner).content
                    if extension == "xlsx":
                        from openpyxl import load_workbook

                        assert (
                            load_workbook(io.BytesIO(saved)).active["A1"].value == "AlpenData synthetic pilot"
                        )
                    else:
                        from pptx import Presentation

                        assert (
                            Presentation(io.BytesIO(saved)).slides[0].shapes.title.text
                            == "AlpenData synthetic pilot"
                        )
                    continue

                def send(message, values):
                    page.locator("iframe").evaluate(
                        "(frame, args) => frame.contentWindow.postMessage("
                        "JSON.stringify({MessageId: args.message, Values: args.values}), args.origin)",
                        {"message": message, "values": values, "origin": qa["office"]},
                    )

                send("Send_UNO_Command", {"Command": ".uno:SelectAll"})
                send(
                    "Action_Paste",
                    {"Mimetype": "text/plain;charset=utf-8", "Data": "Manual revision AlpenData QA"},
                )
                send("Send_UNO_Command", {"Command": ".uno:SelectAll"})
                page.get_by_role(
                    "button",
                    name="Reprendre la sélection avec l’assistant"
                    if language == "fr"
                    else "Work on the selection with the assistant",
                ).click()
                expect(page.locator(".office-selection blockquote")).to_contain_text(
                    "Manual revision", timeout=15000
                )
                save.click()
                file = root + "/files/" + files[extension]
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    versions = client.get(file + "/versions", headers=owner).json()
                    if versions["file"]["version"] >= 2:
                        break
                    time.sleep(0.2)
                assert versions["file"]["version"] == 2
                from docx import Document

                saved = client.get(file + "/versions/2/content", headers=owner).content
                assert "Manual revision" in " ".join(p.text for p in Document(io.BytesIO(saved)).paragraphs)
                result = client.post(
                    file + "/versions",
                    headers=owner,
                    json={
                        "expected_version": 2,
                        "filename": "Pilot.docx",
                        "content_base64": base64.b64encode(
                            office_bytes("docx", "Concurrent assistant version")
                        ).decode(),
                    },
                )
                assert result.status_code == 201
                send(
                    "Action_Paste", {"Mimetype": "text/plain;charset=utf-8", "Data": "Preserve this conflict"}
                )
                save.click()
                deadline = time.monotonic() + 30
                conflict = None
                while time.monotonic() < deadline:
                    with app.state.session_factory() as db:
                        conflict = db.scalar(
                            select(EditorSession.conflict_file_id).where(
                                EditorSession.file_id == files[extension]
                            )
                        )
                    if conflict:
                        break
                    time.sleep(0.2)
                assert conflict
                original = client.get(file + "/versions/3/content", headers=owner).content
                assert "Concurrent assistant" in " ".join(
                    p.text for p in Document(io.BytesIO(original)).paragraphs
                )
                copied = client.get(
                    root + "/files/" + conflict + "/versions/1/content", headers=owner
                ).content
                assert "Preserve this conflict" in " ".join(
                    p.text for p in Document(io.BytesIO(copied)).paragraphs
                )
        finally:
            page.screenshot(path=str(qa["output"] / "last-page.png"))
            (qa["output"] / "last-page.txt").write_text(page.locator("body").inner_text())
            context.close()
            browser.close()
