"""Real isolated Hermes produces, edits, recalculates and renders the four formats."""

import json
import shlex
import shutil
import tempfile
from dataclasses import replace
from pathlib import Path
from uuid import uuid4

import pytest
from model_http import completion, model_http
from test_documents import start_document_turn
from test_routines import connected_service as connected_service
from test_routines import routine_service as routine_service

from alpendata_api.chat_worker import ChatWorker
from alpendata_api.model_gateway import ModelGateway
from alpendata_api.runtime import RuntimeSettings

pytestmark = pytest.mark.linux_only

# This is an actual terminal program run by Hermes, not a mock of document generation.
PROBE = r"""
import json, subprocess, sys, tempfile
from pathlib import Path
sys.path.insert(0, '/opt/hermes/alpendata/runtime')
from document_builder import DocumentSpec, build
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

root = Path('/state/workspace')
brief = DocumentSpec.model_validate({
 'title': 'Préparation de la séance', 'subtitle': 'Coaching individuel',
 'sections': [
  {'heading':'Objectifs', 'paragraphs':[
   'Faire le point sur les priorités de la semaine et les actions engagées.',
   'Identifier les difficultés rencontrées et choisir une prochaine action concrète.']},
  {'heading':'Déroulement', 'columns':['Étape','Durée'], 'rows':[
   ['Retour sur la semaine','15 minutes'],['Travail sur les priorités','30 minutes'],
   ['Prochaine action','15 minutes']]}
 ]})
build(brief, root/'briefing.docx')
build(brief, root/'fiche.pdf')
original = (root/'briefing.docx').read_bytes()
try:
 build(brief, root/'briefing.docx')
 raise AssertionError('Existing file was overwritten')
except ValueError:
 assert (root/'briefing.docx').read_bytes() == original
document = Document(root/'briefing.docx')
assert document.tables[0].cell(1,0).text == 'Retour sur la semaine'
document.add_paragraph('Prochaine séance : vérifier les actions retenues.')
document.save(root/'briefing.docx')
assert Document(root/'briefing.docx').paragraphs[-1].text.startswith('Prochaine séance')

sheet = DocumentSpec.model_validate({'title':'Suivi des séances','sheets':[{
 'name':'Séances','columns':['Activité','Heures'],
 'rows':[['Préparation',1.5],['Coaching',2],['Total',{'formula':'=SUM(B2:B3)','number_format':'0.0'}],
         ['=This is a literal label',None]],'widths':[40,18]}]})
build(sheet, root/'suivi.xlsx')
workbook = load_workbook(root/'suivi.xlsx')
assert workbook['Séances']['A5'].data_type == 's'
workbook['Séances']['B3'] = 2.5
workbook.save(root/'suivi.xlsx')

slides = DocumentSpec.model_validate({'title':'Atelier de coaching','slides':[
 {'title':'Priorités de la semaine','paragraphs':[
 'Quelles actions ont permis de progresser ?',
 'Quelles difficultés demandent encore un accompagnement ?']},
 {'title':'Prochaines actions','paragraphs':[
 'Choisir une action concrète à essayer avant la prochaine séance.',
 'Définir ensemble comment observer les progrès.']} ]})
build(slides, root/'atelier.pptx')
deck = Presentation(root/'atelier.pptx')
assert len(deck.slides) == 2
assert all(shape.has_text_frame for slide in deck.slides for shape in slide.shapes)
deck.core_properties.subject = 'Coaching individuel et en entreprise'
deck.slides[1].shapes[0].text_frame.paragraphs[0].runs[0].text = 'Actions retenues'
deck.save(root/'atelier.pptx')
assert Presentation(root/'atelier.pptx').core_properties.subject.startswith('Coaching')
assert Presentation(root/'atelier.pptx').slides[1].shapes[0].text == 'Actions retenues'
too_long = DocumentSpec.model_validate({'title':'Overflow','slides':[
 {'title':'Title','paragraphs':['Long text ' * 500]}]})
try:
 build(too_long, root/'overflow.pptx')
 raise AssertionError('Overflow was accepted')
except ValueError:
 assert not (root/'overflow.pptx').exists()

rendered = root/'rendered'; rendered.mkdir()
calculated = root/'calculated'; calculated.mkdir()
def convert(source, extension, destination):
 with tempfile.TemporaryDirectory(prefix='alpendata-lo-') as profile:
  subprocess.run(['/usr/bin/libreoffice','-env:UserInstallation='+Path(profile).as_uri(),
   '--headless','--convert-to',extension,'--outdir',str(destination),str(source)],
   timeout=60,check=True,capture_output=True)
 output = destination/(source.stem+'.'+extension)
 assert output.is_file() and output.stat().st_size > 0
 return output

recalculated = convert(root/'suivi.xlsx','xlsx',calculated)
assert load_workbook(recalculated,data_only=True)['Séances']['B4'].value == 4
assert load_workbook(recalculated,data_only=False)['Séances']['B4'].value == '=SUM(B2:B3)'
assert load_workbook(recalculated)['Séances']['A5'].data_type == 's'
pdfs = [root/'fiche.pdf']
for source in (root/'briefing.docx',recalculated,root/'atelier.pptx'):
 pdfs.append(convert(source,'pdf',rendered))
expected = [('Préparation de la séance',1),('Prochaine séance',1),('Total',1),('Actions retenues',2)]
pages = {}
for pdf,(text,count) in zip(pdfs,expected):
 reader = PdfReader(pdf)
 assert len(reader.pages) == count, (pdf.name,len(reader.pages))
 assert text in '\n'.join(page.extract_text() for page in reader.pages), pdf.name
 pages[pdf.name] = len(reader.pages)
 subprocess.run(['/usr/bin/pdftoppm','-scale-to','1400','-png',str(pdf),str(rendered/pdf.stem)],
  check=True,timeout=30,capture_output=True)
print(json.dumps({'verified_pages':pages,'calculated_hours':4,'editable_slides':2}))
"""


def test_document_generation_editing_and_rendering_in_real_hermes(routine_service, request):
    image = request.config.getoption("--runtime-image")
    if not image or not request.config.getoption("--postgresql-bin"):
        pytest.skip("Requires real OCI image and PostgreSQL")
    app, client, settings, _, _, org, alice, bob = routine_service
    path = start_document_turn(routine_service)
    files = ["briefing.docx", "fiche.pdf", "calculated/suivi.xlsx", "atelier.pptx"]
    calls = [
        ("read_file", {"path": "/opt/hermes/alpendata/runtime/DOCUMENT_GUIDE.md"}),
        ("terminal", {"command": "/opt/venv/bin/python -c " + shlex.quote(PROBE), "timeout": 150}),
        *[("alpendata_publish_document", {"path": filename}) for filename in files],
    ]

    def respond(body):
        answers = [item for item in body["messages"] if item["role"] == "tool"]
        if answers:
            assert "Creating and editing workplace documents" in answers[0]["content"]
        if len(answers) > 1:
            terminal = json.loads(answers[1]["content"])
            assert terminal["exit_code"] == 0, terminal
            evidence = json.loads(terminal["output"])
            assert evidence["calculated_hours"] == 4
            assert evidence["editable_slides"] == 2
        if len(answers) == len(calls):
            for answer in answers[2:]:
                assert json.loads(answer["content"])["status"] == 200, answer
            return 200, completion(body, content="Les quatre fichiers sont prêts."), {}
        tool, arguments = calls[len(answers)]
        answer = completion(body, content=None)
        answer["choices"][0]["finish_reason"] = "tool_calls"
        answer["choices"][0]["message"]["tool_calls"] = [
            {
                "id": uuid4().hex[:9],
                "type": "function",
                "function": {"name": tool, "arguments": json.dumps(arguments)},
            }
        ]
        return 200, answer, {}

    with (
        tempfile.TemporaryDirectory(prefix="alpendata-office-") as state,
        model_http(respond) as (transport, _, _),
    ):
        settings = replace(settings, runtime=RuntimeSettings(Path(state), image))
        worker = ChatWorker(
            settings, app.state.session_factory, gateway=ModelGateway(settings.model, session=transport)
        )
        assert worker.run_once()
        turn = client.get(path, headers=bob[2]).json()["turns"][0]
        assert turn["status"] == "completed", turn
        assert {item["filename"] for item in turn["artifacts"]} == {Path(name).name for name in files}
        downloads = {}
        for item in turn["artifacts"]:
            url = f"/api/organizations/{org}/documents/{item['id']}/download"
            response = client.get(url, headers=bob[2])
            assert len(response.content) == item["size"]
            assert client.get(url, headers=alice[2]).status_code == 404
            downloads[item["filename"]] = response.content
        destination = request.config.getoption("--document-qa-output")
        if destination:
            destination = Path(destination)
            destination.mkdir(parents=True, exist_ok=True)
            output = Path(tempfile.mkdtemp(prefix="office-", dir=destination))
            for name, content in downloads.items():
                (output / name).write_bytes(content)
            with worker.runtime.owner_state(org, bob[0]) as owner:
                shutil.copytree(owner / "workspace" / "rendered", output / "rendered")
            print(f"Document visual QA artifacts: {output}")
