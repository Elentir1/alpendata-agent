"""Optional local starter layouts; Hermes can also use the installed libraries directly."""

import argparse
import io
import json
from pathlib import Path
from xml.sax.saxutils import escape

from pydantic import BaseModel, ConfigDict, Field, model_validator


class Input(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)


class Section(Input):
    heading: str = Field(default="", max_length=160)
    paragraphs: list[str] = Field(default_factory=list, max_length=40)
    columns: list[str] = Field(default_factory=list, max_length=6)
    rows: list[list[str]] = Field(default_factory=list, max_length=100)

    @model_validator(mode="after")
    def rectangular(self):
        if self.rows and (
            not self.columns or any(len(row) != len(self.columns) for row in self.rows)
        ):
            raise ValueError("Table rows must match the columns")
        return self


class Cell(Input):
    formula: str = Field(min_length=2, max_length=1024, pattern=r"^=")
    number_format: str = Field(default="General", max_length=100)


class Sheet(Input):
    name: str = Field(min_length=1, max_length=31)
    columns: list[str] = Field(min_length=1, max_length=20)
    rows: list[list[str | int | float | bool | None | Cell]] = Field(
        default_factory=list, max_length=2000
    )
    widths: list[float | int] = Field(default_factory=list, max_length=20)

    @model_validator(mode="after")
    def rectangular(self):
        if any(len(row) != len(self.columns) for row in self.rows):
            raise ValueError("Worksheet rows must match columns")
        if self.widths and (
            len(self.widths) != len(self.columns)
            or any(not 5 <= width <= 80 for width in self.widths)
        ):
            raise ValueError("Specify one width between 5 and 80 for each column")
        return self


class Slide(Input):
    title: str = Field(min_length=1, max_length=160)
    paragraphs: list[str] = Field(default_factory=list, max_length=12)


class DocumentSpec(Input):
    title: str = Field(min_length=1, max_length=200)
    subtitle: str = Field(default="", max_length=500)
    sections: list[Section] = Field(default_factory=list, max_length=80)
    sheets: list[Sheet] = Field(default_factory=list, max_length=30)
    slides: list[Slide] = Field(default_factory=list, max_length=60)


FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
BOLD_FONT = FONT.with_name("DejaVuSans-Bold.ttf")


def word(spec, output):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21), Cm(29.7)
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin = page.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "DejaVu Sans", Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (
        ("Normal", 10.5),
        ("Title", 26),
        ("Heading 1", 15),
        ("Subtitle", 11),
    ):
        style = document.styles[name]
        style.font.name, style.font.size = "DejaVu Sans", Pt(size)
        style.font.color.rgb = RGBColor.from_string("18212F")
        style.font.bold, style.font.italic = name in {"Title", "Heading 1"}, False
        fonts = style.element.rPr.rFonts
        for key in list(fonts.attrib):
            if "theme" in key.lower():
                del fonts.attrib[key]
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), "DejaVu Sans")
        paragraph_properties = style.element.get_or_add_pPr()
        for tag in ("pBdr", "numPr"):
            element = paragraph_properties.find(qn("w:" + tag))
            if element is not None:
                paragraph_properties.remove(element)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(8)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    document.core_properties.author = ""
    document.core_properties.title = spec.title
    document.add_paragraph(spec.title, "Title")
    if spec.subtitle:
        document.add_paragraph(spec.subtitle, "Subtitle")
    for section in spec.sections:
        if section.heading:
            document.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            document.add_paragraph(paragraph)
        if section.columns:
            table = document.add_table(rows=1, cols=len(section.columns))
            table.style = "Normal Table"
            borders = OxmlElement("w:tblBorders")
            for edge in ("insideH", "bottom"):
                line = OxmlElement("w:" + edge)
                for key, value in {
                    "val": "single",
                    "sz": "4",
                    "color": "D6DBE1",
                }.items():
                    line.set(qn("w:" + key), value)
                borders.append(line)
            table._tbl.tblPr.append(borders)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for cell, value in zip(table.rows[0].cells, section.columns):
                cell.text = value
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF1F5")
                cell._tc.get_or_add_tcPr().append(shading)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
            for row in section.rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = value
    document.save(output)


def pdf(spec, output):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        LongTable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        TableStyle,
    )

    pdfmetrics.registerFont(TTFont("AlpenSans", str(FONT)))
    pdfmetrics.registerFont(TTFont("AlpenSansBold", str(BOLD_FONT)))
    body = ParagraphStyle(
        "body", fontName="AlpenSans", fontSize=10.5, leading=15, spaceAfter=8
    )
    heading = ParagraphStyle(
        "heading",
        parent=body,
        fontName="AlpenSansBold",
        fontSize=15,
        leading=20,
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True,
    )
    title = ParagraphStyle(
        "title", parent=heading, fontSize=26, leading=32, spaceBefore=0, spaceAfter=12
    )
    flow = [Paragraph(escape(spec.title), title)]
    if spec.subtitle:
        flow.append(Paragraph(escape(spec.subtitle), body))
    for section in spec.sections:
        if section.heading:
            flow.append(Paragraph(escape(section.heading), heading))
        flow.extend(
            Paragraph(escape(value).replace("\n", "<br/>"), body)
            for value in section.paragraphs
        )
        if section.columns:
            rows = [
                [Paragraph(escape(value), body) for value in row]
                for row in [section.columns, *section.rows]
            ]
            table = LongTable(
                rows,
                colWidths=[470 / len(section.columns)] * len(section.columns),
                repeatRows=1,
                splitByRow=True,
            )
            table.setStyle(
                TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#D6DBE1")),
                ])
            )
            flow.extend([table, Spacer(1, 10)])
    SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=62.5,
        rightMargin=62.5,
        topMargin=56,
        bottomMargin=56,
        title=spec.title,
        author="",
    ).build(flow)


def spreadsheet(spec, output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not spec.sheets:
        raise ValueError("An XLSX needs sheets")
    if len({sheet.name.casefold() for sheet in spec.sheets}) != len(spec.sheets):
        raise ValueError("Worksheet names must be unique")
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator, workbook.properties.title = "", spec.title
    for source in spec.sheets:
        sheet = workbook.create_sheet(source.name)
        # Text values are always text; only an explicit formula object becomes a formula.
        for row_index, values in enumerate([source.columns, *source.rows], 1):
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_index, column)
                if isinstance(value, Cell):
                    cell.value, cell.number_format = value.formula, value.number_format
                else:
                    cell.value = value
                    if isinstance(value, str):
                        cell.data_type = "s"
                cell.font = Font(
                    name="DejaVu Sans", size=11, bold=row_index == 1, color="18212F"
                )
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_index == 1:
                    cell.fill = PatternFill("solid", fgColor="EEF1F5")
        for index in range(1, len(source.columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                source.widths[index - 1] if source.widths else 24
            )
        sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = (
            "landscape" if len(source.columns) > 4 else "portrait"
        )
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    workbook.save(output)


def measured_lines(text, font, width):
    lines = 0
    for paragraph in text.split("\n"):
        used = 0
        lines += 1
        for word in paragraph.split():
            length = font.getlength(word + " ")
            if length > width:
                raise ValueError(
                    "A slide contains an unbreakable word wider than the text area"
                )
            if used and used + length > width:
                lines += 1
                used = 0
            used += length
    return lines


def presentation(spec, output):
    from PIL import ImageFont
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    if not spec.slides:
        raise ValueError("A PPTX needs slides; include a cover only if requested")
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.author, deck.core_properties.title = "", spec.title
    for content in spec.slides:
        title_lines = measured_lines(
            content.title, ImageFont.truetype(str(BOLD_FONT), 32), 810
        )
        body_lines = sum(
            measured_lines(value, ImageFont.truetype(str(FONT), 20), 810)
            for value in content.paragraphs
        )
        if title_lines > 2 or body_lines * 29 + len(content.paragraphs) * 12 > 330:
            raise ValueError(
                "Slide text does not fit; shorten it or explicitly split the slide"
            )
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for top, height, size, bold, paragraphs in (
            (0.55, 1.4, 32, True, [content.title]),
            (2.1, 4.8, 20, False, content.paragraphs),
        ):
            frame = slide.shapes.add_textbox(
                Inches(0.8), Inches(top), Inches(11.7), Inches(height)
            ).text_frame
            frame.word_wrap = True
            frame.margin_left = frame.margin_right = frame.margin_top = (
                frame.margin_bottom
            ) = 0
            for index, text in enumerate(paragraphs):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = text
                paragraph.font.name, paragraph.font.size = "DejaVu Sans", Pt(size)
                paragraph.font.bold = bold
                paragraph.font.color.rgb = RGBColor.from_string("18212F")
                paragraph.space_after = Pt(12)
    deck.save(output)


BUILDERS = {".docx": word, ".pdf": pdf, ".xlsx": spreadsheet, ".pptx": presentation}


def build(spec, output):
    output = Path(output)
    if output.suffix.lower() not in BUILDERS:
        raise ValueError("Choose PDF, DOCX, XLSX or PPTX")
    # Each format must receive its own intended content; no silently ignored sections.
    fields = {
        ".docx": "sections",
        ".pdf": "sections",
        ".xlsx": "sheets",
        ".pptx": "slides",
    }
    wanted = fields[output.suffix.lower()]
    if any(getattr(spec, key) for key in fields.values() if key != wanted):
        raise ValueError("Use only the content field appropriate to this output format")
    if output.exists():
        raise ValueError("Output already exists; choose a new filename")
    content = io.BytesIO()
    BUILDERS[output.suffix.lower()](spec, content)
    if content.tell() > 5 * 1024 * 1024:
        raise ValueError("The document exceeds the 5 MiB download limit")
    # Complete before publishing a path; a failed renderer leaves no partial file.
    with output.open("xb") as destination:
        destination.write(content.getvalue())
    return {"filename": output.name, "size": output.stat().st_size}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spec", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.spec.stat().st_size > 1024 * 1024:
        parser.error("Document specification exceeds 1 MiB")
    spec = DocumentSpec.model_validate_json(args.spec.read_bytes())
    print(json.dumps(build(spec, args.output), ensure_ascii=False))


if __name__ == "__main__":
    main()
