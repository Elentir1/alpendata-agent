"""Untrusted document parsing runs inside a networkless, resource-limited container."""

import base64
import io


def analyze(request):
    data = base64.b64decode(request["content_base64"], validate=True)
    if len(data) > 5 * 1024 * 1024:
        raise ValueError("File too large")
    extension, chunks, remaining = (
        request["filename"].rsplit(".", 1)[-1].lower(),
        [],
        300000,
    )

    partial = False

    def add(reference, text):
        nonlocal remaining, partial
        if remaining <= 0 or len(chunks) >= 500:
            partial = partial or bool(str(text).strip())
            return
        original = str(text).replace("\x00", "")
        text = original[: min(12000, remaining)]
        partial = partial or len(text) < len(original)
        if text.strip():
            chunks.append({"reference": reference, "text": text})
            remaining -= len(text)

    stream = io.BytesIO(data)
    if extension == "pdf":
        from pypdf import PdfReader

        document = PdfReader(stream)
        if document.is_encrypted:
            return {"status": "password_required", "pages": []}
        partial = len(document.pages) > 500
        for index, page in enumerate(document.pages[:500]):
            add(f"page {index + 1}", page.extract_text() or "")
    elif extension == "docx":
        from docx import Document

        document = Document(stream)
        for index, paragraph in enumerate(document.paragraphs):
            add(f"paragraph {index + 1}", paragraph.text)
        for index, table in enumerate(document.tables):
            add(
                f"table {index + 1}",
                "\n".join(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                ),
            )
    elif extension == "xlsx":
        from openpyxl import load_workbook

        document = load_workbook(stream, read_only=True, data_only=True)
        formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        partial = len(document.worksheets) > 100
        try:
            for sheet in document.worksheets[:100]:
                partial = (
                    partial
                    or (sheet.max_row or 0) > 5000
                    or (sheet.max_column or 0) > 100
                )
                formula_rows = formulas[sheet.title].iter_rows(
                    max_row=min(sheet.max_row or 0, 5000),
                    max_col=min(sheet.max_column or 0, 100),
                )
                for index, row in enumerate(
                    sheet.iter_rows(
                        max_row=min(sheet.max_row or 0, 5000),
                        max_col=min(sheet.max_column or 0, 100),
                    ),
                    1,
                ):
                    cells = []
                    for cell, formula in zip(row, next(formula_rows)):
                        if cell.value is not None:
                            cells.append(f"{cell.coordinate}: {cell.value}")
                        elif formula.data_type == "f":
                            cells.append(
                                f"{formula.coordinate}: {formula.value} [formula; result not cached]"
                            )
                    if cells:
                        add(f"{sheet.title}!{index}", " | ".join(cells))
                    if remaining <= 0 or len(chunks) >= 500:
                        partial = True
                        break
        finally:
            document.close()
            formulas.close()
    elif extension == "pptx":
        from pptx import Presentation

        for index, slide in enumerate(Presentation(stream).slides):
            add(
                f"slide {index + 1}",
                "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame),
            )
    elif extension in ("txt", "md", "csv"):
        lines = data.decode("utf-8-sig").splitlines()
        partial = len(lines) > 20000
        for index in range(0, min(len(lines), 20000), 40):
            add(
                f"lines {index + 1}–{min(index + 40, len(lines))}",
                "\n".join(lines[index : index + 40]),
            )
    else:
        return {"status": "vision_required", "pages": []}
    return {
        "status": "ready" if chunks else "no_text",
        "pages": chunks,
        "partial": partial,
    }
